import io
import datetime
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from spark_jobs.sales_analysis import get_filtered_dashboard_data, load_data

# ==========================================
# REPORTLAB NUMBERED CANVAS FOR HEADER/FOOTER
# ==========================================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        
        # We only draw headers/footers on page 1+ if it's a simple multi-page document,
        # but let's make it look premium.
        
        # Color definitions
        muted_color = colors.HexColor("#64748B")
        border_color = colors.HexColor("#E2E8F0")
        
        # Draw header
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#4F46E5"))
        self.drawString(54, 755, "PULSE")
        self.setFont("Helvetica", 8)
        self.setFillColor(muted_color)
        self.drawString(90, 755, "|   Executive Sales Analytics & Performance Report")
        
        # Header line
        self.setStrokeColor(border_color)
        self.setLineWidth(0.5)
        self.line(54, 747, 558, 747)
        
        # Draw footer
        self.line(54, 52, 558, 52)
        self.setFont("Helvetica", 8)
        self.setFillColor(muted_color)
        now_str = datetime.datetime.now().strftime("%B %d, %Y %I:%M %p")
        self.drawString(54, 38, f"Generated: {now_str}")
        self.drawString(250, 38, "Confidential - For Internal Use Only")
        
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 38, page_text)
        
        self.restoreState()


# ==========================================
# REPORT 1: PDF SALES REPORT
# ==========================================
def generate_sales_pdf(dataset_path, year_filter="All", region_filter="All", category_filter="All"):
    # Retrieve data
    data = get_filtered_dashboard_data(
        dataset_path,
        year_filter=year_filter,
        region_filter=region_filter,
        category_filter=category_filter
    )
    
    buffer = io.BytesIO()
    
    # Setup document
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=72,
        bottomMargin=72
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Typography Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=24
    )
    
    h1_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#1E293B"),
        spaceBefore=15,
        spaceAfter=10,
        keepWithNext=True
    )
    
    metric_label_style = ParagraphStyle(
        'MetricLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#4F46E5")
    )
    
    metric_val_style = ParagraphStyle(
        'MetricValue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A")
    )
    
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#334155")
    )
    
    table_text = ParagraphStyle(
        'TableText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1E293B")
    )
    
    table_header = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.white
    )

    story = []
    
    # 1. Header Title & Meta Info
    story.append(Paragraph("Executive Sales Performance", title_style))
    meta_desc = f"Filter Scope: Year = <b>{year_filter}</b> | Region = <b>{region_filter}</b> | Category = <b>{category_filter}</b>"
    story.append(Paragraph(meta_desc, subtitle_style))
    
    # Helper to format values
    def fmt_curr(val):
        is_neg = val < 0
        v = abs(val)
        if v >= 10000000:
            res = f"₹{round(v / 10000000, 2)}Cr"
        elif v >= 100000:
            res = f"₹{round(v / 100000, 2)}L"
        elif v >= 1000:
            res = f"₹{round(v / 1000, 2)}K"
        else:
            res = f"₹{round(v, 2)}"
        return f"-{res}" if is_neg else res

    # 2. KPI Cards Block
    kpi_data = [
        [
            Paragraph("TOTAL SALES", metric_label_style),
            Paragraph("TOTAL PROFIT", metric_label_style),
            Paragraph("TOTAL ORDERS", metric_label_style),
            Paragraph("AVG DISCOUNT", metric_label_style)
        ],
        [
            Paragraph(fmt_curr(data['total_sales']), metric_val_style),
            Paragraph(fmt_curr(data['total_profit']), metric_val_style),
            Paragraph(f"{data['total_orders']:,}", metric_val_style),
            Paragraph(f"{data['avg_discount']}%", metric_val_style)
        ]
    ]
    
    kpi_table = Table(kpi_data, colWidths=[126, 126, 126, 126])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#F1F5F9")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('PADDING', (0,0), (-1,-1), 12),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    
    story.append(kpi_table)
    story.append(Spacer(1, 20))
    
    # 3. Performance Summary Text
    summary_html = f"""
    This executive performance report encapsulates sales analytics computed dynamically. 
    A total of <b>{data['total_orders']:,}</b> orders were processed, driving a total top-line sales volume of 
    <b>{fmt_curr(data['total_sales'])}</b> and yielding an aggregate net profit of <b>{fmt_curr(data['total_profit'])}</b>. 
    The overall average product discount applied across transactions stands at <b>{data['avg_discount']}%</b>.
    """
    story.append(Paragraph(summary_html, body_style))
    story.append(Spacer(1, 15))
    
    # 4. Regional & Segment Performance Section
    story.append(Paragraph("Performance Breakdowns", h1_style))
    
    # Let's organize Regional sales and Segment sales side by side
    region_rows = [[Paragraph("<b>Region</b>", table_text), Paragraph("<b>Sales Vol</b>", table_text)]]
    for r, val in zip(data['regions'], data['region_values']):
        region_rows.append([Paragraph(r, table_text), Paragraph(fmt_curr(val), table_text)])
        
    segment_rows = [[Paragraph("<b>Segment</b>", table_text), Paragraph("<b>Sales Vol</b>", table_text)]]
    for s, val in zip(data['segments'], data['segment_values']):
        segment_rows.append([Paragraph(s, table_text), Paragraph(fmt_curr(val), table_text)])
        
    # Standardize length
    max_len = max(len(region_rows), len(segment_rows))
    while len(region_rows) < max_len:
        region_rows.append(["", ""])
    while len(segment_rows) < max_len:
        segment_rows.append(["", ""])
        
    breakdown_data = []
    for r_row, s_row in zip(region_rows, segment_rows):
        breakdown_data.append([r_row[0], r_row[1], "", s_row[0], s_row[1]])
        
    breakdown_table = Table(breakdown_data, colWidths=[150, 90, 24, 150, 90])
    breakdown_table.setStyle(TableStyle([
        ('LINEBELOW', (0,0), (1,0), 1, colors.HexColor("#4F46E5")),
        ('LINEBELOW', (3,0), (4,0), 1, colors.HexColor("#4F46E5")),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
    ]))
    
    story.append(breakdown_table)
    story.append(Spacer(1, 20))
    
    # 5. Top Performing Products Section
    story.append(Paragraph("Top 5 Performing Products", h1_style))
    
    top_products_data = [
        [
            Paragraph("Product Name", table_header),
            Paragraph("Sales", table_header),
            Paragraph("Profit", table_header)
        ]
    ]
    
    for p in data['top_products']:
        # Format sales and profit for display
        sales_formatted = fmt_curr(p['total_sales_raw'])
        profit_formatted = fmt_curr(p['total_profit_raw'])
        profit_color = "#10B981" if p['total_profit_raw'] >= 0 else "#EF4444"
        profit_p = Paragraph(f"<font color='{profit_color}'><b>{profit_formatted}</b></font>", table_text)
        top_products_data.append([
            Paragraph(p['Product Name'], table_text),
            Paragraph(sales_formatted, table_text),
            profit_p
        ])
        
    products_table = Table(top_products_data, colWidths=[310, 100, 94])
    products_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#4F46E5")),
        ('PADDING', (0,0), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
    ]))
    
    story.append(products_table)
    
    # Build Document
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


# ==========================================
# REPORT 2: EXCEL SALES REPORT (openpyxl)
# ==========================================
def generate_sales_excel(dataset_path, year_filter="All", region_filter="All", category_filter="All"):
    # Retrieve data
    data = get_filtered_dashboard_data(
        dataset_path,
        year_filter=year_filter,
        region_filter=region_filter,
        category_filter=category_filter
    )
    
    wb = openpyxl.Workbook()
    
    # Styles Setup
    title_font = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
    section_font = Font(name='Segoe UI', size=12, bold=True, color='1E293B')
    header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    regular_font = Font(name='Segoe UI', size=10)
    
    indigo_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    light_blue_fill = PatternFill(start_color='F0F9FF', end_color='F0F9FF', fill_type='solid')
    gray_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # ------------------------------------------
    # SHEET 1: Summary Dashboard
    # ------------------------------------------
    ws1 = wb.active
    ws1.title = "Summary Dashboard"
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Banner
    ws1.merge_cells('A1:D2')
    ws1['A1'] = "PULSE SALES EXECUTIVE SUMMARY"
    ws1['A1'].font = title_font
    ws1['A1'].fill = indigo_fill
    ws1['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Meta Details
    ws1['A4'] = "Scope:"
    ws1['A4'].font = bold_font
    ws1['B4'] = f"Year: {year_filter} | Region: {region_filter} | Category: {category_filter}"
    ws1['B4'].font = regular_font
    ws1['A5'] = "Generated:"
    ws1['A5'].font = bold_font
    ws1['B5'] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws1['B5'].font = regular_font
    
    # KPI Grid
    kpi_headers = ["Metric", "Value"]
    kpi_values = [
        ("Total Sales", data['total_sales']),
        ("Total Profit", data['total_profit']),
        ("Total Orders", data['total_orders']),
        ("Average Discount", data['avg_discount'] / 100.0)
    ]
    
    ws1['A7'] = "Key Performance Indicators (KPIs)"
    ws1['A7'].font = section_font
    
    for c_idx, h_text in enumerate(kpi_headers, start=1):
        cell = ws1.cell(row=8, column=c_idx, value=h_text)
        cell.font = header_font
        cell.fill = indigo_fill
        cell.alignment = Alignment(horizontal='center')
        
    for r_idx, (m_title, m_val) in enumerate(kpi_values, start=9):
        c1 = ws1.cell(row=r_idx, column=1, value=m_title)
        c2 = ws1.cell(row=r_idx, column=2, value=m_val)
        
        c1.font = bold_font
        c1.fill = gray_fill
        c1.border = thin_border
        
        c2.font = regular_font
        c2.border = thin_border
        
        if "Sales" in m_title or "Profit" in m_title:
            c2.number_format = '₹#,##0.00'
        elif "Orders" in m_title:
            c2.number_format = '#,##0'
        elif "Discount" in m_title:
            c2.number_format = '0.0%'
            
    # Top Products Grid
    ws1['A15'] = "Top 5 Performing Products"
    ws1['A15'].font = section_font
    
    prod_headers = ["Product Name", "Sales Volume", "Net Profit"]
    for c_idx, h_text in enumerate(prod_headers, start=1):
        cell = ws1.cell(row=16, column=c_idx, value=h_text)
        cell.font = header_font
        cell.fill = indigo_fill
        cell.alignment = Alignment(horizontal='center')
        
    for r_idx, p in enumerate(data['top_products'], start=17):
        c1 = ws1.cell(row=r_idx, column=1, value=p['Product Name'])
        c2 = ws1.cell(row=r_idx, column=2, value=p['total_sales_raw'])
        c3 = ws1.cell(row=r_idx, column=3, value=p['total_profit_raw'])
        
        for c in [c1, c2, c3]:
            c.font = regular_font
            c.border = thin_border
            
        c2.number_format = '₹#,##0.00'
        c3.number_format = '₹#,##0.00'
        
    # Auto-adjust column widths for Sheet 1
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Avoid using merged cell value lengths to distort single column sizing
            if cell.coordinate in ['A1', 'B1', 'C1', 'D1', 'A2', 'B2', 'C2', 'D2']:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)
        
    # ------------------------------------------
    # SHEET 2: Monthly & Category Analysis
    # ------------------------------------------
    ws2 = wb.create_sheet(title="Trends & Distributions")
    ws2.views.sheetView[0].showGridLines = True
    
    # Monthly Trends
    ws2['A1'] = "Monthly Sales Breakdown"
    ws2['A1'].font = section_font
    
    ws2['A2'] = "Month"
    ws2['B2'] = "Sales Volume"
    ws2['A2'].font = header_font
    ws2['A2'].fill = indigo_fill
    ws2['B2'].font = header_font
    ws2['B2'].fill = indigo_fill
    
    for r_idx, (m, val) in enumerate(zip(data['months'], data['month_values']), start=3):
        c1 = ws2.cell(row=r_idx, column=1, value=m)
        c2 = ws2.cell(row=r_idx, column=2, value=val)
        c1.font = regular_font
        c1.border = thin_border
        c2.font = regular_font
        c2.border = thin_border
        c2.number_format = '₹#,##0.00'
        
    # Category Distribution
    ws2['D1'] = "Category Sales Breakdown"
    ws2['D1'].font = section_font
    
    ws2['D2'] = "Category"
    ws2['E2'] = "Sales Volume"
    ws2['D2'].font = header_font
    ws2['D2'].fill = indigo_fill
    ws2['E2'].font = header_font
    ws2['E2'].fill = indigo_fill
    
    for r_idx, (cat, val) in enumerate(zip(data['categories'], data['category_values']), start=3):
        c1 = ws2.cell(row=r_idx, column=4, value=cat)
        c2 = ws2.cell(row=r_idx, column=5, value=val)
        c1.font = regular_font
        c1.border = thin_border
        c2.font = regular_font
        c2.border = thin_border
        c2.number_format = '₹#,##0.00'
        
    # Region Distribution
    ws2['G1'] = "Regional Sales Breakdown"
    ws2['G1'].font = section_font
    
    ws2['G2'] = "Region"
    ws2['H2'] = "Sales Volume"
    ws2['G2'].font = header_font
    ws2['G2'].fill = indigo_fill
    ws2['H2'].font = header_font
    ws2['H2'].fill = indigo_fill
    
    for r_idx, (reg, val) in enumerate(zip(data['regions'], data['region_values']), start=3):
        c1 = ws2.cell(row=r_idx, column=7, value=reg)
        c2 = ws2.cell(row=r_idx, column=8, value=val)
        c1.font = regular_font
        c1.border = thin_border
        c2.font = regular_font
        c2.border = thin_border
        c2.number_format = '₹#,##0.00'
        
    # Auto-adjust column widths for Sheet 2
    for col_col in ['A', 'B', 'D', 'E', 'G', 'H']:
        max_len = 0
        for cell in ws2[col_col]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws2.column_dimensions[col_col].width = max(max_len + 4, 14)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ==========================================
# REPORT 3: ML PREDICTION PDF REPORT (Matplotlib + ReportLab)
# ==========================================
def generate_prediction_pdf(dataset_path):
    # Retrieve dynamic data (contains forecast info)
    data = get_filtered_dashboard_data(dataset_path)
    fc = data['forecast']
    
    # 1. Generate Matplotlib Chart
    plt.figure(figsize=(7, 3.5))
    
    hist_labels = fc['historical_labels'][-12:] # Show last 12 months for clarity
    hist_values = fc['historical_values'][-12:]
    
    forecast_labels = fc['forecast_labels']
    forecast_values = fc['forecast_values']
    
    # We stitch the last historical point with the first forecast point
    stitch_labels = [hist_labels[-1]] + forecast_labels if hist_labels else forecast_labels
    stitch_values = [hist_values[-1]] + forecast_values if hist_values else forecast_values
    
    # Format labels to fit cleanly
    # Plot historical
    plt.plot(hist_labels, [v/1000 for v in hist_values], marker='o', color='#4F46E5', linewidth=2.5, label='Historical Sales')
    
    # Plot forecast
    plt.plot(stitch_labels, [v/1000 for v in stitch_values], marker='s', linestyle='--', color='#0EA5E9', linewidth=2, label='Forecast Projection')
    
    plt.title('Sales Trend & Machine Learning Forecast (₹ in Thousands)', fontsize=11, fontweight='bold', color='#1E293B', pad=12)
    plt.ylabel('Sales Volume (₹K)', fontsize=9, color='#64748B')
    plt.grid(True, linestyle=':', alpha=0.6, color='#CBD5E1')
    
    plt.xticks(rotation=20, ha='right', fontsize=8, color='#64748B')
    plt.yticks(fontsize=8, color='#64748B')
    plt.legend(frameon=True, facecolor='#F8FAFC', edgecolor='#E2E8F0', fontsize=8)
    
    # Style boundaries
    ax = plt.gca()
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#E2E8F0')
    ax.spines['bottom'].set_color('#E2E8F0')
    
    plt.tight_layout()
    
    chart_buffer = io.BytesIO()
    plt.savefig(chart_buffer, format='png', dpi=200)
    plt.close()
    chart_buffer.seek(0)
    
    # 2. Build ReportLab Doc
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=72,
        bottomMargin=72
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=20
    )
    
    h1_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#1E293B"),
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#334155")
    )
    
    table_text = ParagraphStyle(
        'TableText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1E293B")
    )
    
    table_header = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.white
    )
    
    story = []
    
    # Title
    story.append(Paragraph("Predictive Analytics Report", title_style))
    story.append(Paragraph("Machine Learning Sales Forecasting & Trend Projections", subtitle_style))
    
    # Metric Callout Table
    def fmt_val(val):
        if val >= 100000:
            return f"₹{round(val / 100000, 2)}L"
        elif val >= 1000:
            return f"₹{round(val / 1000, 2)}K"
        return f"₹{round(val, 2)}"
        
    trend_prefix = "+" if fc['forecast_trend'] >= 0 else ""
    trend_color = "#10B981" if fc['forecast_trend'] >= 0 else "#EF4444"
    
    kpi_data = [
        [
            Paragraph("FORECAST NEXT MONTH", ParagraphStyle('L1', parent=metric_label_style, fontSize=8)),
            Paragraph("PROJECTION TREND", ParagraphStyle('L2', parent=metric_label_style, fontSize=8)),
            Paragraph("MODEL CONFIDENCE / R² ACCURACY", ParagraphStyle('L3', parent=metric_label_style, fontSize=8))
        ],
        [
            Paragraph(fmt_val(fc['forecast_value']), metric_val_style),
            Paragraph(f"<font color='{trend_color}'><b>{trend_prefix}{fc['forecast_trend']}%</b></font>", metric_val_style),
            Paragraph(f"{fc['forecast_accuracy']}%", metric_val_style)
        ]
    ]
    
    kpi_table = Table(kpi_data, colWidths=[168, 168, 168])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#F1F5F9")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('PADDING', (0,0), (-1,-1), 10),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    
    story.append(kpi_table)
    story.append(Spacer(1, 15))
    
    # Insight block
    insight_desc = f"""
    The forecasting algorithm utilizes a <b>Linear Regression Model</b> trained on chronological aggregated monthly sales. 
    The current model calibration yields a confidence metric of <b>{fc['forecast_accuracy']}%</b>. 
    Based on historical sales signals, next month's sales are projected to reach <b>{fmt_val(fc['forecast_value'])}</b>, representing a 
    <b>{trend_prefix}{fc['forecast_trend']}%</b> shift compared to the immediate prior month.
    """
    story.append(Paragraph(insight_desc, body_style))
    story.append(Spacer(1, 15))
    
    # Visual Chart
    img_width = 460
    img_height = 230
    chart_flowable = Image(chart_buffer, width=img_width, height=img_height)
    story.append(KeepTogether([
        Paragraph("Forecast Visual Model", h1_style),
        chart_flowable
    ]))
    story.append(Spacer(1, 15))
    
    # Forecast Table
    fc_table_data = [
        [
            Paragraph("Projection Period", table_header),
            Paragraph("Forecasted Sales", table_header),
            Paragraph("Status / Recommendation", table_header)
        ]
    ]
    
    rec_texts = [
        "Normal operation. Review inventory limits.",
        "Expected period transition. Stock alignment recommended.",
        "Extended projection. Strategy plan check."
    ]
    
    for idx, (label, val) in enumerate(zip(forecast_labels, forecast_values)):
        story_val_fmt = fmt_val(val)
        rec = rec_texts[idx] if idx < len(rec_texts) else "Strategy plan check."
        fc_table_data.append([
            Paragraph(label, table_text),
            Paragraph(story_val_fmt, table_text),
            Paragraph(rec, table_text)
        ])
        
    fc_table = Table(fc_table_data, colWidths=[120, 130, 254])
    fc_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#4F46E5")),
        ('PADDING', (0,0), (-1,-1), 6),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('LINEBELOW', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
    ]))
    
    story.append(KeepTogether([
        Paragraph("Forecast Summary Matrix", h1_style),
        fc_table
    ]))
    
    # Build Document
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


# ==========================================
# REPORT 4: DYNAMIC SPARK DATASET SUMMARY STATISTICS
# ==========================================
def get_spark_dataset_summary(dataset_path):
    df = load_data(dataset_path)
    total_rows = df.count()
    cols = df.columns
    
    # Compute schema info
    schema_details = []
    numeric_cols = []
    for field in df.schema.fields:
        schema_details.append({
            "name": field.name,
            "type": str(field.dataType)
        })
        if str(field.dataType).startswith("DoubleType") or str(field.dataType).startswith("IntegerType") or str(field.dataType).startswith("LongType") or str(field.dataType).startswith("FloatType"):
            numeric_cols.append(field.name)
            
    # Calculate null counts dynamically
    null_counts = {}
    from pyspark.sql.functions import sum as spark_sum, col
    null_exprs = [spark_sum(col(c).isNull().cast("int")).alias(c) for c in cols]
    try:
        null_res = df.agg(*null_exprs).collect()[0].asDict()
        for c in cols:
            null_counts[c] = null_res.get(c, 0)
    except Exception as e:
        # Fallback if any error occurs
        for c in cols:
            null_counts[c] = 0

    # Calculate Descriptive Statistics using Spark describe()
    desc_stats = []
    if numeric_cols:
        try:
            # describe() computes count, mean, stddev, min, max
            desc_df = df.describe(*numeric_cols).collect()
            # Convert row list into a dict of stats
            # Row(summary='count', Sales='9994', Profit='9994')
            stats_by_summary = {}
            for row in desc_df:
                row_dict = row.asDict()
                summary_type = row_dict.pop('summary')
                stats_by_summary[summary_type] = row_dict
                
            for col_name in numeric_cols:
                count_val = stats_by_summary.get('count', {}).get(col_name, "0")
                mean_val = stats_by_summary.get('mean', {}).get(col_name, "0.0")
                std_val = stats_by_summary.get('stddev', {}).get(col_name, "0.0")
                min_val = stats_by_summary.get('min', {}).get(col_name, "0.0")
                max_val = stats_by_summary.get('max', {}).get(col_name, "0.0")
                
                # Format floats cleanly
                try:
                    mean_val = f"{float(mean_val):,.2f}"
                except:
                    pass
                try:
                    std_val = f"{float(std_val):,.2f}"
                except:
                    pass
                try:
                    min_val = f"{float(min_val):,.2f}"
                except:
                    pass
                try:
                    max_val = f"{float(max_val):,.2f}"
                except:
                    pass
                
                desc_stats.append({
                    "column": col_name,
                    "count": count_val,
                    "mean": mean_val,
                    "stddev": std_val,
                    "min": min_val,
                    "max": max_val,
                    "null_count": null_counts.get(col_name, 0)
                })
        except Exception as e:
            print("Spark statistics error:", e)
            
    # If no descriptive stats succeeded, fill placeholder
    if not desc_stats:
        for c in cols:
            desc_stats.append({
                "column": c,
                "count": str(total_rows),
                "mean": "N/A",
                "stddev": "N/A",
                "min": "N/A",
                "max": "N/A",
                "null_count": null_counts.get(c, 0)
            })

    return {
        "success": True,
        "rows": total_rows,
        "cols_count": len(cols),
        "columns": cols,
        "schema": schema_details,
        "desc_stats": desc_stats,
        "null_counts": null_counts
    }
